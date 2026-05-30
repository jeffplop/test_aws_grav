import os
import sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

# Asegurar que el path local esté disponible
local_path = os.path.join(os.environ['USERPROFILE'], 'AppData', 'Roaming', 'Python', 'Python313', 'site-packages')
if local_path not in sys.path:
    sys.path.append(local_path)

def add_slide(prs, title, text_content=None, image_path=None):
    # Usar layout 5 (Title Only) o 1 (Title and Content)
    layout = prs.slide_layouts[1] if text_content else prs.slide_layouts[5]
    slide = prs.slides.add_slide(layout)
    
    # Estilizar Título
    title_shape = slide.shapes.title
    title_shape.text = title
    for paragraph in title_shape.text_frame.paragraphs:
        paragraph.font.name = 'Arial'
        paragraph.font.size = Pt(36)
        paragraph.font.bold = True
        paragraph.font.color.rgb = RGBColor(35, 47, 62) # AWS Dark Blue
        
    if text_content:
        body = slide.placeholders[1]
        tf = body.text_frame
        tf.clear() # Limpiar bullet points default
        for line in text_content.split('\n'):
            p = tf.add_paragraph()
            p.text = line
            p.font.size = Pt(20)
            p.font.name = 'Arial'
            p.space_after = Pt(14)
            p.font.color.rgb = RGBColor(60, 60, 60)

    if image_path and os.path.exists(image_path):
        # Insertar imagen centrada o al lado
        left = Inches(4.5) if text_content else Inches(1)
        top = Inches(2.5) if text_content else Inches(1.8)
        height = Inches(4)
        try:
            slide.shapes.add_picture(image_path, left, top, height=height)
        except Exception as e:
            print(f"No se pudo incrustar la imagen {image_path}: {e}")

def create_modern_pptx():
    prs = Presentation()
    
    # 0. Portada
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    title = slide.shapes.title
    subtitle = slide.placeholders[1]
    title.text = "Innovatech Chile: Proyecto Semestral DevOps"
    subtitle.text = "Evaluación Parcial 2 - Defensa Técnica\nDespliegue Automatizado en AWS"
    
    for paragraph in title.text_frame.paragraphs:
        paragraph.font.bold = True
        paragraph.font.color.rgb = RGBColor(255, 153, 0) # AWS Orange

    # 1. Introducción y Requerimientos (Contexto)
    add_slide(prs, "Contexto del Proyecto: Innovatech Chile", 
              "Objetivo: Desplegar la aplicación 'Tienda de Perritos' en AWS.\n"
              "Requerimientos Clave:\n"
              "1. Contenedorización de Frontend y Backend.\n"
              "2. Persistencia de datos crítica.\n"
              "3. Pipeline CI/CD en GitHub Actions.\n"
              "4. Funcionamiento y Comunicación Segura en EC2.")

    # 2. Diseño de Contenedorización (IE1 & IE6)
    add_slide(prs, "[IE1/IE6] Diseño de Contenedorización", 
              "Backend (Node.js):\n"
              "• Multi-stage build: Reduce tamaño aislando dependencias.\n"
              "• Seguridad: Ejecutado con usuario 'node' (non-root).\n\n"
              "Frontend (Nginx):\n"
              "• Multi-stage build: Compilación y Servidor.\n"
              "• Seguridad: nginxinc/nginx-unprivileged (Puerto 8080).\n"
              "• Optimización: Uso de imágenes Alpine Linux.")

    # 3. Orquestación y Requerimientos No Funcionales (IE2 & IE6)
    add_slide(prs, "[IE2/IE6] Docker Compose y Persistencia", 
              "Red Interna: 'app-network' (Bridge DNS discovery).\n"
              "Dependencias Seguras: Backend no inicia sin 'service_healthy' de DB.\n"
              "Persistencia de Datos:\n"
              "• Elección: Named Volume ('db_data').\n"
              "• Justificación: Los Named Volumes aíslan los datos del host,\n"
              "  mejorando la portabilidad y seguridad vs Bind Mounts.\n"
              "• Garantiza la continuidad operativa ante fallos.")

    # 4. Pipeline CI/CD: Flujo y Secrets (IE3)
    add_slide(prs, "[IE3] Pipeline CI/CD en GitHub Actions", 
              "Triggers: Activación exclusiva en la rama 'deploy'.\n"
              "Gestión de Secrets: \n"
              "• Credenciales de AWS y Tokens protegidos.\n"
              "Fases del Pipeline:\n"
              "1. Build Docker (Multi-stage).\n"
              "2. Push: Publicación en Amazon ECR.\n"
              "3. Deploy: Actualización automática en Amazon EC2.",
              "Entrega-Final-EP2/evidencias/diagramas/2-flujo-cicd.png")

    # 5. Justificación Técnica CI/CD (IE7)
    add_slide(prs, "[IE7] Justificación Técnica del Pipeline", 
              "¿Por qué Amazon ECR?\n"
              "• Integración nativa con IAM y menor latencia hacia EC2.\n"
              "¿Por qué este diseño automatiza la entrega continua?\n"
              "• Elimina la intervención humana post-commit.\n"
              "• AWS Systems Manager (SSM) despliega sin requerir abrir \n  el puerto 22 (SSH), mitigando riesgos críticos para Innovatech.",
              "Entrega-Final-EP2/evidencias/screenshots/3-github-actions.png")

    # 6. Principios DevOps Aplicados (IE8)
    add_slide(prs, "[IE8] Principios DevOps Aplicados", 
              "A. Contenedorización: Docker resuelve el 'en mi máquina funciona'.\n"
              "B. Gestión de Entornos: Configuración inyectada vía variables.\n"
              "C. CI/CD: Automatización reduce el 'Lead Time for Changes'.\n"
              "D. Control de Versiones: Git centralizado.\n"
              "E. Persistencia Segura: Volúmenes externos al ciclo efímero.\n"
              "Impacto: Favorece la escalabilidad y mantenibilidad en AWS.")

    # 7. Evidencia: Funcionamiento en EC2 (IE4 & IE9)
    add_slide(prs, "[IE4/IE9] Funcionamiento en EC2", 
              "El contenedor Frontend opera públicamente respondiendo peticiones.\n"
              "El Backend y Base de Datos están restringidos por Security Groups.\n"
              "Los endpoints son consumidos exitosamente por la SPA estática.",
              "Entrega-Final-EP2/evidencias/screenshots/5-app-frontend.png")

    # 8. Evidencia: Comunicación y Contenedores (IE4 & IE5)
    add_slide(prs, "Evidencia Técnica: Contenedores Vivos", 
              "Verificación mediante 'docker ps' y 'docker volume ls'.\n"
              "Todos los contenedores muestran estado '(healthy)'.",
              "Entrega-Final-EP2/evidencias/screenshots/4-docker-ps.png")
              
    # 9. Conclusión y Demo
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    title = slide.shapes.title
    title.text = "Demostración Funcional en Vivo (IE9)"
    body = slide.placeholders[1]
    tf = body.text_frame
    tf.text = "1. Navegación al Frontend en IP Pública.\n2. Inserción de productos (Verificación de Backend).\n3. Eliminación de Contenedores y Recuperación (Verificación de Persistencia).\n4. Despliegue automático tras Git Push."

    out_path = "Entrega-Final-EP2/Presentacion-Semestral-EP2.pptx"
    prs.save(out_path)
    print(f"Generado exitosamente: {out_path}")

if __name__ == "__main__":
    create_modern_pptx()
